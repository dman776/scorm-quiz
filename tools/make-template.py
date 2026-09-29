#!/usr/bin/env python3
"""Generate template.xlsx for the SCORM Quiz Builder Excel import.

Three sheets:
  Instructions - how to fill the template (per question type)
  Settings     - assessment-level key/value settings
  Questions    - one row per question, with example rows for every type

The Questions example rows form a valid mini quiz so they round-trip through the
importer (see tests/xlsx.test.js).
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "template.xlsx"

NAVY = "0B5CAD"
LIGHT = "EEF4FB"
GREY = "F2F4F7"
BORDER = Border(*(Side(style="thin", color="C3C9D0"),) * 4)
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color=NAVY)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = Workbook()

# ----------------------------------------------------------------- Instructions
ins = wb.active
ins.title = "Instructions"
ins.sheet_view.showGridLines = False
ins["A1"] = "SCORM Quiz Builder - Excel Import Template"
ins["A1"].font = TITLE_FONT
lines = [
    "",
    "Fill in the Settings and Questions sheets, then import this file in the SCORM Quiz Builder (Import Excel).",
    "",
    "QUESTIONS SHEET COLUMNS",
    "  ID                 Optional. Leave blank to auto-generate from the prompt.",
    "  Type               One of: single_select, single_select_pill, multiple_select, multiple_select_pill,",
    "                     true_false, single_checkbox, matching, sequence, numeric, short_answer, drag_drop.",
    "  Prompt             The question text shown to the learner. (Required)",
    "  Points             Point value for the question. Defaults to 1.",
    "  Objective          Optional learning objective.",
    "  Section            Optional grouping/topic.",
    "  Scoring            Optional: all_or_nothing, partial, or weighted (choice/matching/sequence/drag_drop).",
    "  Options            The answers. Encoding depends on Type (see below).",
    "  Correct            Used for numeric (the value), short_answer (accepted answers), true_false, single_checkbox.",
    "  Tolerance          Numeric only. Allowed +/- range. Defaults to 0.",
    "  Units              Numeric only. Optional label shown after the input.",
    "  Correct Feedback   Optional feedback shown when the learner is correct.",
    "  Incorrect Feedback Optional feedback shown when the learner is incorrect.",
    "  Rationale          Optional explanation shown in review.",
    "",
    "OPTIONS ENCODING BY TYPE",
    "  single_select / multiple_select (and _pill variants):",
    "     Put every choice in Options separated by a pipe |  and mark correct choices with a leading *",
    "     Example:  *Central switch or hub | Closed loop of neighbors | Single backbone cable",
    "     Optional per-answer score in brackets:  *HTTPS[1.5] | *SSH[1.5] | HTTP[-1] | Telnet[-1]",
    "  true_false:",
    "     Leave Options blank. Put True or False in the Correct column.",
    "  single_checkbox (acknowledgement):",
    "     Options = the statement text. Correct = true (default) or false.",
    "  matching:",
    "     Options = pairs separated by |  using =  ->  HTTPS=443 | SSH=22 | DNS=53 | HTTP=80",
    "  sequence (ordering):",
    "     Options = items in the CORRECT order separated by >  ->  Physical > Data Link > Network > Transport",
    "  numeric:",
    "     Leave Options blank. Correct = the value (e.g. 254). Tolerance and Units optional.",
    "  short_answer:",
    "     Leave Options blank. Correct = accepted answers separated by |  ->  NAT | Network Address Translation",
    "  drag_drop (items into labeled boxes):",
    "     Options = item=zone pairs separated by |  ->  Router=Network | Switch=Data Link | Repeater=",
    "     Zones are created in the order they first appear. Item=Zone A ; Zone B accepts either zone.",
    "     Nothing after = makes a distractor (correct to leave unplaced). =Zone adds a zone with no correct item.",
    "     Zone[n] limits a zone to n items, e.g. Router=Network[1]. Background-image zones need the builder.",
    "",
    "TIP: The example rows already in the Questions sheet are a working quiz. Edit or delete them.",
]
for i, text in enumerate(lines, start=2):
    c = ins.cell(row=i, column=1, value=text)
    c.font = Font(name="Arial", bold=text.isupper() and text.strip() != "", size=10,
                  color=NAVY if text.isupper() and text.strip() != "" else "1A1A1A")
ins.column_dimensions["A"].width = 118

# --------------------------------------------------------------------- Settings
st = wb.create_sheet("Settings")
st.sheet_view.showGridLines = False
st["A1"] = "Setting"
st["B1"] = "Value"
for col in ("A", "B"):
    st[f"{col}1"].font = HEAD_FONT
    st[f"{col}1"].fill = PatternFill("solid", fgColor=NAVY)
    st[f"{col}1"].border = BORDER
settings_rows = [
    ("Title", "ASCEND Sample Quiz"),
    ("Description", "A sample quiz created from the Excel template."),
    ("Version", "1.0"),
    ("Author", "Your Name"),
    ("Language", "en-US"),
    ("Passing Percent", 80),
    ("Max Attempts", 2),
    ("Score Retention", "highest"),
    ("Shuffle Questions", "no"),
    ("Shuffle Answers", "no"),
    ("Show Correct Answers", "yes"),
    ("Delay Answers Until Final Attempt", "yes"),
    ("Pass Message", "Congratulations, you passed."),
    ("Fail Message", "You did not reach the passing score. Please review and retry."),
]
for r, (k, v) in enumerate(settings_rows, start=2):
    a = st.cell(row=r, column=1, value=k)
    b = st.cell(row=r, column=2, value=v)
    a.font = Font(name="Arial", bold=True, size=10)
    b.font = BODY_FONT
    b.fill = PatternFill("solid", fgColor="FFFDE7")  # editable = pale yellow
    for cell in (a, b):
        cell.border = BORDER
        cell.alignment = WRAP
st.column_dimensions["A"].width = 32
st.column_dimensions["B"].width = 60

# -------------------------------------------------------------------- Questions
qs = wb.create_sheet("Questions")
qs.sheet_view.showGridLines = False
headers = ["ID", "Type", "Prompt", "Points", "Objective", "Section", "Scoring",
           "Options", "Correct", "Tolerance", "Units",
           "Correct Feedback", "Incorrect Feedback", "Rationale"]
for c, name in enumerate(headers, start=1):
    cell = qs.cell(row=1, column=c, value=name)
    cell.font = HEAD_FONT
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER

examples = [
    ["", "single_select", "In a star topology, how are devices connected?", 2,
     "Identify network topologies", "Topologies", "",
     "*Each device connects to a central switch or hub | Each device connects to two neighbors in a loop | All devices share one backbone cable",
     "", "", "", "Correct, the central node is the single connection point.",
     "Review the topology comparison module.", "A star centralizes connectivity through one switch or hub."],
    ["", "single_select_pill", "Which medium is least susceptible to electromagnetic interference?", 2,
     "Select transmission media", "Media", "",
     "Unshielded twisted pair | Coaxial cable | *Fiber optic | Powerline",
     "", "", "", "Yes, fiber uses light and is immune to EMI.",
     "Copper media are affected by EMI; fiber is not.", "Fiber transmits light, not electrical signals."],
    ["", "multiple_select", "Select ALL layers that belong to the OSI model.", 4,
     "Classify OSI layers", "OSI Model", "partial",
     "*Physical | *Transport | *Application | Cloud | Virtualization",
     "", "", "", "All three OSI layers identified.",
     "Revisit the seven OSI layers.", "Physical, Transport, and Application are OSI layers."],
    ["", "multiple_select_pill", "Select the protocols that provide encryption.", 3,
     "Recognize secure protocols", "Security", "weighted",
     "*HTTPS[1.5] | *SSH[1.5] | HTTP[-1] | Telnet[-1]",
     "", "", "", "Both encrypted protocols selected.",
     "HTTPS and SSH are the encrypted options.", "HTTPS and SSH encrypt traffic; HTTP and Telnet do not."],
    ["", "true_false", "Power over Ethernet can deliver both data and power over one cable.", 1,
     "Explain PoE", "Power", "", "", "True", "", "",
     "Correct, PoE carries data and power together.",
     "PoE does carry both data and power.", "PoE injects DC power onto the data twisted pair."],
    ["", "single_checkbox", "Acknowledgement statement (learner checks to agree).", 1,
     "Acknowledge policy", "Policy", "",
     "I acknowledge I have read the network safety policy.", "true", "", "",
     "Thank you for acknowledging.", "You must acknowledge to continue.",
     "Acknowledgement is recorded as a true/false interaction."],
    ["", "matching", "Match each protocol to its default port.", 4,
     "Map protocols to ports", "Protocols", "partial",
     "HTTPS=443 | SSH=22 | DNS=53 | HTTP=80", "", "", "",
     "All ports matched correctly.", "Review well-known port numbers.",
     "HTTPS 443, SSH 22, DNS 53, HTTP 80."],
    ["", "sequence", "Arrange the OSI layers from Layer 1 to Layer 4.", 3,
     "Order OSI layers", "OSI Model", "all_or_nothing",
     "Physical > Data Link > Network > Transport", "", "", "",
     "Perfect ordering.", "The order is Physical, Data Link, Network, Transport.",
     "OSI layers 1 to 4 in order."],
    ["", "numeric", "How many usable host addresses are in a /24 IPv4 subnet?", 2,
     "Compute host counts", "Addressing", "", "", "254", "0", "hosts",
     "Correct, 256 minus network and broadcast equals 254.",
     "A /24 has 256 addresses, 254 usable.", "2^8 = 256, minus network and broadcast."],
    ["", "short_answer", "What acronym describes translating private IPs to a public IP?", 1,
     "Recall address translation", "Addressing", "", "",
     "NAT | Network Address Translation", "", "",
     "Correct, NAT translates private to public addressing.",
     "The answer is NAT.", "NAT lets many private hosts share public addresses."],
    ["", "drag_drop", "Drag each device onto the OSI layer where it primarily operates.", 3,
     "Map devices to OSI layers", "OSI Model", "partial",
     "Router=Network | Switch=Data Link | Bridge=Data Link | Repeater=", "", "", "",
     "Correct placement.", "Check which address each device reads: IP or MAC.",
     "Routers read IP (Layer 3); switches and bridges read MAC (Layer 2); a repeater is Layer 1."],
]
for r, row in enumerate(examples, start=2):
    for c, val in enumerate(row, start=1):
        cell = qs.cell(row=r, column=c, value=val)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER
        if r % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=GREY)

widths = [10, 20, 46, 8, 24, 14, 16, 52, 22, 10, 10, 30, 30, 40]
for c, w in enumerate(widths, start=1):
    qs.column_dimensions[qs.cell(row=1, column=c).column_letter].width = w
qs.freeze_panes = "A2"

# Type dropdown for guidance (non-blocking)
types = "single_select,single_select_pill,multiple_select,multiple_select_pill,true_false,single_checkbox,matching,sequence,numeric,short_answer,drag_drop"
dv = DataValidation(type="list", formula1=f'"{types}"', allow_blank=True, showDropDown=False)
qs.add_data_validation(dv)
dv.add(f"B2:B200")

wb.save(OUT)
print(f"Wrote {OUT}")
